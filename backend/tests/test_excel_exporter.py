import io
import openpyxl
from app.models import CalculationParameters, Branch, Cable
from app.routing_engine import solve_routing_and_sizing
from app.excel_exporter import generate_excel_report
from app.sample_generator import generate_sample_excel_workbook


def test_excel_export_generation():
    params = CalculationParameters()
    branches = [
        Branch(branch_id="BR_01", node_from="A", node_to="B", level="Level 1", length_m=10.0),
        Branch(branch_id="RISER_01", node_from="B", node_to="C", level="Transition", branch_type="vertical", length_m=4.5),
    ]
    cables = [
        Cable(cable_tag="C1", source_node="A", dest_node="C", cable_type="power", od_mm=20.0, count=2),
        Cable(cable_tag="C2", source_node="A", dest_node="C", cable_type="control", od_mm=12.0, count=4),
    ]
    calc_resp = solve_routing_and_sizing(params, branches, cables)

    excel_stream = generate_excel_report(calc_resp)
    assert isinstance(excel_stream, io.BytesIO)
    content = excel_stream.getvalue()
    assert len(content) > 1000

    # Read back with openpyxl to verify valid workbook and expected sheets
    wb = openpyxl.load_workbook(io.BytesIO(content))
    expected_sheets = ["Executive Summary", "Tray Sizing Results", "Cable Schedule", "Diagnostics", "Bill of Materials"]
    assert wb.sheetnames == expected_sheets

    # Verify content in Tray Sizing Results
    ws_branches = wb["Tray Sizing Results"]
    assert ws_branches.max_row >= 3  # Header + 2 branches
    assert ws_branches.cell(row=2, column=1).value == "BR_01"

    # Verify Bill of Materials sheet content
    ws_bom = wb["Bill of Materials"]
    assert ws_bom.max_row >= 5
    assert "Bill of Materials" in str(ws_bom.cell(row=1, column=1).value)


def test_sample_template_generation():
    template_stream = generate_sample_excel_workbook()
    assert isinstance(template_stream, io.BytesIO)
    content = template_stream.getvalue()
    assert len(content) > 1000

    wb = openpyxl.load_workbook(io.BytesIO(content))
    assert "Cables" in wb.sheetnames
    assert "Branches" in wb.sheetnames
    assert wb["Cables"].max_row > 10
    assert wb["Branches"].max_row > 10
