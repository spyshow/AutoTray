import io
from datetime import datetime
from typing import Dict, Any, Union
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from .models import CalculationResponse


def generate_excel_report(data: Union[CalculationResponse, Dict[str, Any]]) -> io.BytesIO:
    if isinstance(data, dict):
        response = CalculationResponse.model_validate(data)
    else:
        response = data

    wb = openpyxl.Workbook()

    # Define Styles
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=16, bold=True, color="0F172A")
    subtitle_font = Font(name="Calibri", size=10, italic=True, color="64748B")
    section_font = Font(name="Calibri", size=12, bold=True, color="1E293B")
    kpi_val_font = Font(name="Calibri", size=18, bold=True, color="0F172A")
    kpi_lbl_font = Font(name="Calibri", size=9, color="64748B", bold=True)
    regular_font = Font(name="Calibri", size=10, color="1E293B")
    bold_regular_font = Font(name="Calibri", size=10, bold=True, color="1E293B")

    # Status Fills
    status_ok_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    status_ok_font = Font(name="Calibri", size=10, bold=True, color="166534")

    status_overfill_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    status_overfill_font = Font(name="Calibri", size=10, bold=True, color="991B1B")

    status_empty_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    status_empty_font = Font(name="Calibri", size=10, italic=True, color="64748B")

    stripe_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )
    kpi_card_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_align = Alignment(horizontal="left", vertical="center")
    right_align = Alignment(horizontal="right", vertical="center")

    # ==========================================
    # SHEET 1: Summary & KPIs
    # ==========================================
    ws_summary = wb.active
    ws_summary.title = "Executive Summary"
    ws_summary.views.sheetView[0].showGridLines = True

    # Title Banner
    ws_summary["A1"] = "AutoTray-Router: Cable Tray Sizing & Routing Report"
    ws_summary["A1"].font = title_font
    ws_summary["A2"] = f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')} | Multi-Level Industrial Engine"
    ws_summary["A2"].font = subtitle_font

    # KPI Summary Cards Block
    kpi_items = [
        ("Total Cables Routed", f"{response.summary.total_cables_routed}"),
        ("Unrouted Cables", f"{len(response.summary.unrouted_cables)}"),
        ("Total Branches", f"{response.summary.total_branches}"),
        ("Max Fill Ratio", f"{response.summary.max_fill_pct}%"),
        ("Overfilled Trays", f"{response.summary.overfilled_branches_count}"),
        ("Total Cable Length", f"{response.summary.total_cable_length_routed_m} m"),
        ("Total Tray Network", f"{response.summary.total_tray_length_m} m"),
    ]

    row_kpi_lbl = 4
    row_kpi_val = 5
    for idx, (label, val) in enumerate(kpi_items):
        col = idx + 1
        ws_summary.cell(row=row_kpi_lbl, column=col, value=label).font = kpi_lbl_font
        ws_summary.cell(row=row_kpi_lbl, column=col).alignment = center_align
        ws_summary.cell(row=row_kpi_lbl, column=col).fill = kpi_card_fill
        ws_summary.cell(row=row_kpi_lbl, column=col).border = thin_border

        c_val = ws_summary.cell(row=row_kpi_val, column=col, value=val)
        c_val.font = kpi_val_font
        c_val.alignment = center_align
        c_val.fill = kpi_card_fill
        c_val.border = thin_border

    # Highlights Table
    ws_summary.cell(row=8, column=1, value="System Highlights & Network Diagnostics").font = section_font
    highlights = [
        ("Max Fill Branch ID", response.summary.max_fill_branch_id or "N/A"),
        ("Max Fill Ratio", f"{response.summary.max_fill_pct}%"),
        ("Overfilled Branches (>900mm)", str(response.summary.overfilled_branches_count)),
        ("Disconnected Nodes", ", ".join(response.diagnostics.disconnected_nodes) or "None"),
        ("Missing Referenced Nodes", ", ".join(response.diagnostics.missing_nodes_referenced_in_cables) or "None"),
        ("Unrouted Cables Count", str(len(response.summary.unrouted_cables))),
    ]
    cur_row = 10
    for key, val in highlights:
        c1 = ws_summary.cell(row=cur_row, column=1, value=key)
        c1.font = bold_regular_font
        c1.border = thin_border
        c2 = ws_summary.cell(row=cur_row, column=2, value=val)
        c2.font = regular_font
        c2.border = thin_border
        cur_row += 1

    # ==========================================
    # SHEET 2: Branch Sizing Results
    # ==========================================
    ws_branches = wb.create_sheet(title="Tray Sizing Results")
    ws_branches.views.sheetView[0].showGridLines = True

    branch_headers = [
        "Branch ID",
        "From Node",
        "To Node",
        "Level",
        "Orientation",
        "Length (m)",
        "Height (mm)",
        "Total Cables",
        "Power Cables",
        "Ctrl Cables",
        "Data Cables",
        "Power Width (mm)",
        "Ctrl Width (mm)",
        "Divider (mm)",
        "Calc Width (mm)",
        "Commercial Width (mm)",
        "Fill Ratio (%)",
        "Status",
        "Cable Load (kg/m)",
        "Tray Dead Load (kg/m)",
        "Design Load (kg/m)",
        "Support Span (m)",
        "Supports Count",
        "Mounting Style",
        "Routed Cables List",
    ]

    for col_num, h in enumerate(branch_headers, 1):
        cell = ws_branches.cell(row=1, column=col_num, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border
    ws_branches.row_dimensions[1].height = 28

    for r_idx, b in enumerate(response.branches, 2):
        m_str = "Wall Cantilever" if getattr(b, "support_mounting_type", "ceiling_trapeze") == "wall_cantilever" else "Ceiling Trapeze"
        row_cells = [
            ws_branches.cell(row=r_idx, column=1, value=b.branch_id),
            ws_branches.cell(row=r_idx, column=2, value=b.node_from),
            ws_branches.cell(row=r_idx, column=3, value=b.node_to),
            ws_branches.cell(row=r_idx, column=4, value=b.level),
            ws_branches.cell(row=r_idx, column=5, value=str(b.branch_type or "").capitalize()),
            ws_branches.cell(row=r_idx, column=6, value=b.length_m),
            ws_branches.cell(row=r_idx, column=7, value=b.tray_height_mm),
            ws_branches.cell(row=r_idx, column=8, value=b.cable_count),
            ws_branches.cell(row=r_idx, column=9, value=b.power_cables_count),
            ws_branches.cell(row=r_idx, column=10, value=b.control_cables_count),
            ws_branches.cell(row=r_idx, column=11, value=b.data_cables_count),
            ws_branches.cell(row=r_idx, column=12, value=b.power_width_mm),
            ws_branches.cell(row=r_idx, column=13, value=b.control_width_mm),
            ws_branches.cell(row=r_idx, column=14, value=b.barrier_width_mm),
            ws_branches.cell(row=r_idx, column=15, value=b.calculated_width_mm),
            ws_branches.cell(row=r_idx, column=16, value=b.recommended_commercial_width_mm),
            ws_branches.cell(row=r_idx, column=17, value=b.fill_ratio_pct),
            ws_branches.cell(row=r_idx, column=18, value=b.status),
            ws_branches.cell(row=r_idx, column=19, value=getattr(b, "cable_load_kg_m", 0.0)),
            ws_branches.cell(row=r_idx, column=20, value=getattr(b, "tray_dead_load_kg_m", 0.0)),
            ws_branches.cell(row=r_idx, column=21, value=getattr(b, "total_load_kg_m", 0.0)),
            ws_branches.cell(row=r_idx, column=22, value=getattr(b, "recommended_support_span_m", 2.0)),
            ws_branches.cell(row=r_idx, column=23, value=getattr(b, "supports_count", 1)),
            ws_branches.cell(row=r_idx, column=24, value=m_str),
            ws_branches.cell(row=r_idx, column=25, value=", ".join(str(x) for x in b.cables_routed)),
        ]

        is_stripe = (r_idx % 2 == 1)
        for c in row_cells:
            c.font = regular_font
            c.border = thin_border
            if is_stripe:
                c.fill = stripe_fill

        # Alignments
        row_cells[0].alignment = left_align
        row_cells[1].alignment = left_align
        row_cells[2].alignment = left_align
        row_cells[3].alignment = center_align
        row_cells[4].alignment = center_align
        for idx in range(5, 17):
            row_cells[idx].alignment = right_align
        row_cells[17].alignment = center_align
        for idx in range(18, 23):
            row_cells[idx].alignment = right_align
        row_cells[23].alignment = center_align
        row_cells[24].alignment = left_align

        # Status badge coloring
        status_cell = row_cells[17]
        if b.status == "OK":
            status_cell.fill = status_ok_fill
            status_cell.font = status_ok_font
        elif b.status == "OVERFILL_SPLIT_TIER":
            status_cell.fill = status_overfill_fill
            status_cell.font = status_overfill_font
        else:
            status_cell.fill = status_empty_fill
            status_cell.font = status_empty_font

        ws_branches.row_dimensions[r_idx].height = 20

    # ==========================================
    # SHEET 3: Cable Schedule & Routing
    # ==========================================
    ws_cables = wb.create_sheet(title="Cable Schedule")
    ws_cables.views.sheetView[0].showGridLines = True

    cable_headers = [
        "Cable Tag",
        "Source Node",
        "Dest Node",
        "Type",
        "OD (mm)",
        "Count",
        "Routing Status",
        "Total Length (m)",
        "Path Traversed",
        "Unrouted Reason",
    ]

    for col_num, h in enumerate(cable_headers, 1):
        cell = ws_cables.cell(row=1, column=col_num, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border
    ws_cables.row_dimensions[1].height = 28

    for r_idx, c in enumerate(response.cables, 2):
        path_str = " ➔ ".join(c.path_nodes) if c.path_nodes else "None"
        row_cells = [
            ws_cables.cell(row=r_idx, column=1, value=c.cable_tag),
            ws_cables.cell(row=r_idx, column=2, value=c.source_node),
            ws_cables.cell(row=r_idx, column=3, value=c.dest_node),
            ws_cables.cell(row=r_idx, column=4, value=str(c.cable_type or "").upper()),
            ws_cables.cell(row=r_idx, column=5, value=c.od_mm),
            ws_cables.cell(row=r_idx, column=6, value=c.count),
            ws_cables.cell(row=r_idx, column=7, value=c.status),
            ws_cables.cell(row=r_idx, column=8, value=c.total_length_m),
            ws_cables.cell(row=r_idx, column=9, value=path_str),
            ws_cables.cell(row=r_idx, column=10, value=c.unrouted_reason or ""),
        ]

        is_stripe = (r_idx % 2 == 1)
        for cell in row_cells:
            cell.font = regular_font
            cell.border = thin_border
            if is_stripe:
                cell.fill = stripe_fill

        row_cells[0].alignment = left_align
        row_cells[1].alignment = left_align
        row_cells[2].alignment = left_align
        row_cells[3].alignment = center_align
        row_cells[4].alignment = right_align
        row_cells[5].alignment = right_align
        row_cells[6].alignment = center_align
        row_cells[7].alignment = right_align
        row_cells[8].alignment = left_align
        row_cells[9].alignment = left_align

        if c.status == "ROUTED":
            row_cells[6].fill = status_ok_fill
            row_cells[6].font = status_ok_font
        else:
            row_cells[6].fill = status_overfill_fill
            row_cells[6].font = status_overfill_font

        ws_cables.row_dimensions[r_idx].height = 20

    # ==========================================
    # SHEET 4: Diagnostics
    # ==========================================
    ws_diag = wb.create_sheet(title="Diagnostics")
    ws_diag.views.sheetView[0].showGridLines = True

    diag_headers = ["Issue Type", "Item Tag / Identifier", "Details / Reason"]
    for col_num, h in enumerate(diag_headers, 1):
        cell = ws_diag.cell(row=1, column=col_num, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border
    ws_diag.row_dimensions[1].height = 28

    d_row = 2
    for node in response.diagnostics.disconnected_nodes:
        ws_diag.cell(row=d_row, column=1, value="Disconnected Node").font = regular_font
        ws_diag.cell(row=d_row, column=2, value=node).font = regular_font
        ws_diag.cell(row=d_row, column=3, value="Node has no connections or is isolated in a separate network component").font = regular_font
        for c in range(1, 4):
            ws_diag.cell(row=d_row, column=c).border = thin_border
        d_row += 1

    for node in response.diagnostics.missing_nodes_referenced_in_cables:
        ws_diag.cell(row=d_row, column=1, value="Missing Node in Cable").font = regular_font
        ws_diag.cell(row=d_row, column=2, value=node).font = regular_font
        ws_diag.cell(row=d_row, column=3, value="Cable source or dest node is not defined in any branch").font = regular_font
        for c in range(1, 4):
            ws_diag.cell(row=d_row, column=c).border = thin_border
        d_row += 1

    for unrouted in response.diagnostics.unrouted_cables_details:
        ws_diag.cell(row=d_row, column=1, value="Unrouted Cable").font = regular_font
        ws_diag.cell(row=d_row, column=2, value=unrouted.cable_tag).font = regular_font
        ws_diag.cell(row=d_row, column=3, value=unrouted.unrouted_reason or "No path found").font = regular_font
        for c in range(1, 4):
            ws_diag.cell(row=d_row, column=c).border = thin_border
        d_row += 1

    if d_row == 2:
        # No issues found
        ws_diag.cell(row=2, column=1, value="None").font = status_ok_font
        ws_diag.cell(row=2, column=2, value="Network Fully Connected").font = status_ok_font
        ws_diag.cell(row=2, column=3, value="All cables routed successfully without topology errors.").font = status_ok_font
        for c in range(1, 4):
            ws_diag.cell(row=2, column=c).border = thin_border

    # ==========================================
    # SHEET 5: Bill of Materials (BOM)
    # ==========================================
    ws_bom = wb.create_sheet(title="Bill of Materials")
    ws_bom.views.sheetView[0].showGridLines = True

    # Title Banner
    ws_bom["A1"] = "Bill of Materials (BOM) & Material Take-Off"
    ws_bom["A1"].font = title_font
    ws_bom["A2"] = f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')} | AutoTray-Router Engineering MTO"
    ws_bom["A2"].font = subtitle_font

    cur_row = 4

    # 1. Trays Table
    ws_bom.cell(row=cur_row, column=1, value="1. CABLE TRAY & RISER MATERIAL SCHEDULE").font = section_font
    cur_row += 1

    tray_headers = ["Commercial Width (mm)", "Side Height (mm)", "Orientation", "Total Length (m)", "Standard 3m Sections (pcs)", "Branch Segments"]
    for col_num, h in enumerate(tray_headers, 1):
        cell = ws_bom.cell(row=cur_row, column=col_num, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border
    ws_bom.row_dimensions[cur_row].height = 26
    cur_row += 1

    bom_data = response.bom
    if bom_data and bom_data.trays:
        for t in bom_data.trays:
            ws_bom.cell(row=cur_row, column=1, value=t.width_mm).font = bold_regular_font
            ws_bom.cell(row=cur_row, column=2, value=t.height_mm).font = regular_font
            ws_bom.cell(row=cur_row, column=3, value=t.branch_type.capitalize()).font = regular_font
            ws_bom.cell(row=cur_row, column=4, value=t.total_length_m).font = regular_font
            ws_bom.cell(row=cur_row, column=5, value=t.section_count_3m).font = bold_regular_font
            ws_bom.cell(row=cur_row, column=6, value=t.branch_count).font = regular_font
            for c in range(1, 7):
                ws_bom.cell(row=cur_row, column=c).border = thin_border
                ws_bom.cell(row=cur_row, column=c).alignment = center_align
            cur_row += 1

        # Total Trays Row
        ws_bom.cell(row=cur_row, column=1, value="Total Trays").font = bold_regular_font
        ws_bom.cell(row=cur_row, column=4, value=bom_data.total_tray_length_m).font = bold_regular_font
        ws_bom.cell(row=cur_row, column=5, value=bom_data.total_sections_3m).font = bold_regular_font
        for c in range(1, 7):
            cell = ws_bom.cell(row=cur_row, column=c)
            cell.border = thin_border
            cell.fill = stripe_fill
            cell.alignment = center_align
        cur_row += 2
    else:
        ws_bom.cell(row=cur_row, column=1, value="No tray materials routed.").font = status_empty_font
        cur_row += 2

    # 2. Cable Tray Fittings & In-Line Reducers Schedule
    ws_bom.cell(row=cur_row, column=1, value="2. CABLE TRAY FITTINGS & IN-LINE REDUCERS SCHEDULE").font = section_font
    cur_row += 1

    # 2A. Fittings Table
    fit_headers = ["Fitting Description / Item", "Nominal Width (mm)", "Side Height (mm)", "Quantity (pcs)", "Applicable Junction Nodes"]
    for col_num, h in enumerate(fit_headers, 1):
        cell = ws_bom.cell(row=cur_row, column=col_num, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border
    ws_bom.row_dimensions[cur_row].height = 26
    cur_row += 1

    if bom_data and bom_data.fittings:
        for fit in bom_data.fittings:
            ws_bom.cell(row=cur_row, column=1, value=fit.fitting_name).font = bold_regular_font
            ws_bom.cell(row=cur_row, column=2, value=fit.width_mm).font = regular_font
            ws_bom.cell(row=cur_row, column=3, value=fit.height_mm).font = regular_font
            ws_bom.cell(row=cur_row, column=4, value=fit.quantity).font = bold_regular_font
            ws_bom.cell(row=cur_row, column=5, value=", ".join(fit.nodes)).font = regular_font
            for c in range(1, 6):
                ws_bom.cell(row=cur_row, column=c).border = thin_border
                ws_bom.cell(row=cur_row, column=c).alignment = center_align if c in [2, 3, 4] else left_align
            cur_row += 1
        cur_row += 1
    else:
        ws_bom.cell(row=cur_row, column=1, value="No tray fittings required.").font = status_empty_font
        cur_row += 2

    # 2B. Reducers Table
    red_headers = ["Reduction Step (W1 -> W2)", "From Width (mm)", "To Width (mm)", "Side Height (mm)", "Geometry Type", "Quantity (pcs)", "Installed Locations (Node:Branch)"]
    for col_num, h in enumerate(red_headers, 1):
        cell = ws_bom.cell(row=cur_row, column=col_num, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border
    ws_bom.row_dimensions[cur_row].height = 26
    cur_row += 1

    if bom_data and bom_data.reducers:
        for red in bom_data.reducers:
            ws_bom.cell(row=cur_row, column=1, value=f"{red.from_width_mm}mm -> {red.to_width_mm}mm").font = bold_regular_font
            ws_bom.cell(row=cur_row, column=2, value=red.from_width_mm).font = regular_font
            ws_bom.cell(row=cur_row, column=3, value=red.to_width_mm).font = regular_font
            ws_bom.cell(row=cur_row, column=4, value=red.height_mm).font = regular_font
            ws_bom.cell(row=cur_row, column=5, value=red.reducer_type.replace('_', ' ').title()).font = regular_font
            ws_bom.cell(row=cur_row, column=6, value=red.quantity).font = bold_regular_font
            loc_str = ", ".join(f"{loc.get('node_id', '')}:{loc.get('branch_id', '')}" for loc in red.locations)
            ws_bom.cell(row=cur_row, column=7, value=loc_str).font = regular_font
            for c in range(1, 8):
                ws_bom.cell(row=cur_row, column=c).border = thin_border
                ws_bom.cell(row=cur_row, column=c).alignment = center_align if c in [2, 3, 4, 5, 6] else left_align
            cur_row += 1
        cur_row += 1
    else:
        ws_bom.cell(row=cur_row, column=1, value="No in-line reducers required.").font = status_empty_font
        cur_row += 2

    # 3. Accessories Table
    ws_bom.cell(row=cur_row, column=1, value="3. INSTALLATION ACCESSORIES & HARDWARE").font = section_font
    cur_row += 1

    acc_headers = ["Item Name", "Category", "Specification / Description", "Estimated Quantity", "Unit"]
    for col_num, h in enumerate(acc_headers, 1):
        cell = ws_bom.cell(row=cur_row, column=col_num, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border
    ws_bom.row_dimensions[cur_row].height = 26
    cur_row += 1

    if bom_data and bom_data.accessories:
        for acc in bom_data.accessories:
            ws_bom.cell(row=cur_row, column=1, value=acc.item_name).font = bold_regular_font
            ws_bom.cell(row=cur_row, column=2, value=acc.category).font = regular_font
            ws_bom.cell(row=cur_row, column=3, value=acc.description).font = regular_font
            ws_bom.cell(row=cur_row, column=4, value=acc.quantity).font = bold_regular_font
            ws_bom.cell(row=cur_row, column=5, value=acc.unit).font = regular_font
            for c in range(1, 6):
                ws_bom.cell(row=cur_row, column=c).border = thin_border
                if c in [2, 4, 5]:
                    ws_bom.cell(row=cur_row, column=c).alignment = center_align
                else:
                    ws_bom.cell(row=cur_row, column=c).alignment = left_align
            cur_row += 1
        cur_row += 1
    else:
        ws_bom.cell(row=cur_row, column=1, value="No accessories required.").font = status_empty_font
        cur_row += 2

    # 4. Cable Length Summary Table
    ws_bom.cell(row=cur_row, column=1, value="4. CABLE SCHEDULE LENGTH TAKE-OFF").font = section_font
    cur_row += 1

    cable_bom_headers = ["Cable Category / Type", "Routed Cable Count", "Total Routed Distance (m)", "Average Run (m)"]
    for col_num, h in enumerate(cable_bom_headers, 1):
        cell = ws_bom.cell(row=cur_row, column=col_num, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border
    ws_bom.row_dimensions[cur_row].height = 26
    cur_row += 1

    if bom_data and bom_data.cables_summary:
        for csum in bom_data.cables_summary:
            ws_bom.cell(row=cur_row, column=1, value=csum.cable_type.upper()).font = bold_regular_font
            ws_bom.cell(row=cur_row, column=2, value=csum.cable_count).font = regular_font
            ws_bom.cell(row=cur_row, column=3, value=csum.total_routed_length_m).font = bold_regular_font
            ws_bom.cell(row=cur_row, column=4, value=csum.avg_length_m).font = regular_font
            for c in range(1, 5):
                ws_bom.cell(row=cur_row, column=c).border = thin_border
                ws_bom.cell(row=cur_row, column=c).alignment = center_align
            cur_row += 1

        # Total Cables Row
        ws_bom.cell(row=cur_row, column=1, value="Total All Routed Cables").font = bold_regular_font
        ws_bom.cell(row=cur_row, column=3, value=bom_data.total_cable_length_m).font = bold_regular_font
        for c in range(1, 5):
            cell = ws_bom.cell(row=cur_row, column=c)
            cell.border = thin_border
            cell.fill = stripe_fill
            cell.alignment = center_align
        cur_row += 1
    else:
        ws_bom.cell(row=cur_row, column=1, value="No routed cables found.").font = status_empty_font
        cur_row += 2

    # 5. Structural Loading & Support Schedule
    cur_row += 1
    ws_bom.cell(row=cur_row, column=1, value="5. STRUCTURAL WEIGHT & SUPPORT SCHEDULE (IEC 61537 / NEMA VE 1)").font = section_font
    cur_row += 1

    struct_headers = ["Structural Metric", "Calculated Value", "Unit", "Engineering Standard"]
    for col_num, h in enumerate(struct_headers, 1):
        cell = ws_bom.cell(row=cur_row, column=col_num, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border
    ws_bom.row_dimensions[cur_row].height = 26
    cur_row += 1

    struct_rows = [
        ("Total Routed Cable Weight", getattr(bom_data, "total_cable_weight_kg", 0.0) if bom_data else 0.0, "kg", "Technical handbook / Catalog / IEC density"),
        ("Cable Tray Steel Dead Weight", getattr(bom_data, "total_tray_weight_kg", 0.0) if bom_data else 0.0, "kg", "Sheet steel @ 7850 kg/m³, 15% perforation discount"),
        ("Total Installation Structural Weight", getattr(bom_data, "total_installation_weight_kg", 0.0) if bom_data else 0.0, "kg", "Cable weight + Tray dead weight"),
        ("Total Required Support Locations", getattr(bom_data, "total_supports_count", 0) if bom_data else 0, "locations", "IEC 61537 SWL safe spans + NEMA VE-2 junction supports"),
    ]
    for metric, val, unit, ref in struct_rows:
        ws_bom.cell(row=cur_row, column=1, value=metric).font = bold_regular_font
        ws_bom.cell(row=cur_row, column=2, value=val).font = bold_regular_font
        ws_bom.cell(row=cur_row, column=3, value=unit).font = regular_font
        ws_bom.cell(row=cur_row, column=4, value=ref).font = regular_font
        for c in range(1, 5):
            ws_bom.cell(row=cur_row, column=c).border = thin_border
            if c in [2, 3]:
                ws_bom.cell(row=cur_row, column=c).alignment = center_align
            else:
                ws_bom.cell(row=cur_row, column=c).alignment = left_align
        cur_row += 1
    cur_row += 1

    # =========================================================================
    # SHEET 6: Fittings & Reducers (Dedicated Sheet)
    # =========================================================================
    ws_fittings = wb.create_sheet(title="Fittings & Reducers")
    ws_fittings.views.sheetView[0].showGridLines = True

    ws_fittings.cell(row=2, column=1, value="AUTO-TRAY ROUTER - FITTINGS & REDUCERS SCHEDULE").font = title_font
    ws_fittings.cell(row=3, column=1, value="Comprehensive junction fittings schedule, nominal dimensions, and port-specific reduction take-off.").font = subtitle_font

    r_cur = 5
    ws_fittings.cell(row=r_cur, column=1, value="1. TRAY FITTINGS SCHEDULE").font = section_font
    r_cur += 1

    f_headers = ["Fitting Item Name", "Fitting Type", "Nominal Width (mm)", "Side Height (mm)", "Quantity (pcs)", "Applicable Nodes"]
    for col_num, h in enumerate(f_headers, 1):
        cell = ws_fittings.cell(row=r_cur, column=col_num, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border
    ws_fittings.row_dimensions[r_cur].height = 26
    r_cur += 1

    if bom_data and bom_data.fittings:
        for fit in bom_data.fittings:
            ws_fittings.cell(row=r_cur, column=1, value=fit.fitting_name).font = bold_regular_font
            ws_fittings.cell(row=r_cur, column=2, value=fit.fitting_type).font = regular_font
            ws_fittings.cell(row=r_cur, column=3, value=fit.width_mm).font = regular_font
            ws_fittings.cell(row=r_cur, column=4, value=fit.height_mm).font = regular_font
            ws_fittings.cell(row=r_cur, column=5, value=fit.quantity).font = bold_regular_font
            ws_fittings.cell(row=r_cur, column=6, value=", ".join(fit.nodes)).font = regular_font
            for c in range(1, 7):
                ws_fittings.cell(row=r_cur, column=c).border = thin_border
                ws_fittings.cell(row=r_cur, column=c).alignment = center_align if c in [2, 3, 4, 5] else left_align
            r_cur += 1
        r_cur += 1
    else:
        ws_fittings.cell(row=r_cur, column=1, value="No tray fittings required.").font = status_empty_font
        r_cur += 2

    # 2. Reducers Schedule
    ws_fittings.cell(row=r_cur, column=1, value="2. IN-LINE TRAY REDUCERS SCHEDULE").font = section_font
    r_cur += 1

    for col_num, h in enumerate(red_headers, 1):
        cell = ws_fittings.cell(row=r_cur, column=col_num, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border
    ws_fittings.row_dimensions[r_cur].height = 26
    r_cur += 1

    if bom_data and bom_data.reducers:
        for red in bom_data.reducers:
            ws_fittings.cell(row=r_cur, column=1, value=f"{red.from_width_mm}mm -> {red.to_width_mm}mm").font = bold_regular_font
            ws_fittings.cell(row=r_cur, column=2, value=red.from_width_mm).font = regular_font
            ws_fittings.cell(row=r_cur, column=3, value=red.to_width_mm).font = regular_font
            ws_fittings.cell(row=r_cur, column=4, value=red.height_mm).font = regular_font
            ws_fittings.cell(row=r_cur, column=5, value=red.reducer_type.replace('_', ' ').title()).font = regular_font
            ws_fittings.cell(row=r_cur, column=6, value=red.quantity).font = bold_regular_font
            loc_str = ", ".join(f"{loc.get('node_id', '')}:{loc.get('branch_id', '')}" for loc in red.locations)
            ws_fittings.cell(row=r_cur, column=7, value=loc_str).font = regular_font
            for c in range(1, 8):
                ws_fittings.cell(row=r_cur, column=c).border = thin_border
                ws_fittings.cell(row=r_cur, column=c).alignment = center_align if c in [2, 3, 4, 5, 6] else left_align
            r_cur += 1
        r_cur += 1
    else:
        ws_fittings.cell(row=r_cur, column=1, value="No in-line reducers required.").font = status_empty_font
        r_cur += 2

    # 3. Node Configuration Detail
    if response.nodes:
        ws_fittings.cell(row=r_cur, column=1, value="3. NETWORK JUNCTION NODES TOPOLOGY").font = section_font
        r_cur += 1
        n_headers = ["Node Tag", "Level / Elevation", "Connected Branches", "Selected Fitting Type", "Nominal Dimensions (WxH)", "Reducers Required"]
        for col_num, h in enumerate(n_headers, 1):
            cell = ws_fittings.cell(row=r_cur, column=col_num, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align
            cell.border = thin_border
        ws_fittings.row_dimensions[r_cur].height = 26
        r_cur += 1

        for node in response.nodes:
            b_list = ", ".join(f"{b.branch_id}({b.width_mm}mm)" for b in node.connected_branches)
            red_count = len([r for r in node.reducers.values() if r.enabled])
            ws_fittings.cell(row=r_cur, column=1, value=node.node_id).font = bold_regular_font
            ws_fittings.cell(row=r_cur, column=2, value=node.level).font = regular_font
            ws_fittings.cell(row=r_cur, column=3, value=b_list).font = regular_font
            ws_fittings.cell(row=r_cur, column=4, value=node.selected_fitting_type).font = regular_font
            ws_fittings.cell(row=r_cur, column=5, value=f"{node.width_mm}x{node.height_mm} mm").font = bold_regular_font
            ws_fittings.cell(row=r_cur, column=6, value=f"{red_count} active").font = regular_font
            for c in range(1, 7):
                ws_fittings.cell(row=r_cur, column=c).border = thin_border
                ws_fittings.cell(row=r_cur, column=c).alignment = center_align if c in [2, 4, 5, 6] else left_align
            r_cur += 1

    # Auto-adjust column widths for all sheets (capped to 50 for readability)
    for ws in [ws_summary, ws_branches, ws_cables, ws_diag, ws_bom, ws_fittings]:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                if "\n" in val_str:
                    lines = val_str.split("\n")
                    max_len = max(max_len, max(len(l) for l in lines))
                else:
                    max_len = max(max_len, len(val_str))
            ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 50)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
