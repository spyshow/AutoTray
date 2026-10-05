import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def generate_sample_excel_workbook() -> io.BytesIO:
    wb = openpyxl.Workbook()

    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    regular_font = Font(name="Calibri", size=10, color="000000")
    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    right_align = Alignment(horizontal="right", vertical="center")

    # Sheet 1: Cables
    ws_cables = wb.active
    ws_cables.title = "Cables"
    ws_cables.views.sheetView[0].showGridLines = True

    cable_headers = ["Tag", "Source Node", "Destination Node", "Type", "OD (mm)", "Count"]
    for col_idx, h in enumerate(cable_headers, 1):
        c = ws_cables.cell(row=1, column=col_idx, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = center_align
        c.border = thin_border
    ws_cables.row_dimensions[1].height = 25

    sample_cables = [
        ("C_PWR_01", "MCC_L1", "FOREHEARTH_FAN_M1", "Power", 28.4, 1),
        ("C_PWR_02", "MCC_L1", "COOLING_PUMP_P1", "Power", 24.2, 1),
        ("C_PWR_03", "MCC_L1", "FEEDER_PANEL_L2", "Power", 35.0, 1),
        ("C_PWR_04", "TRANSF_01", "MCC_L1", "Power", 42.5, 3),
        ("C_PWR_05", "MCC_L2", "MIXER_MOTOR_M2", "Power", 22.8, 1),
        ("C_PWR_06", "MCC_L2", "EXHAUST_BLOWER_B1", "Power", 31.0, 1),
        ("C_PWR_07", "MCC_L1", "COMPRESSOR_CP1", "Power", 38.5, 1),
        ("C_CTRL_01", "DCS_RACK_L1", "FOREHEARTH_FAN_M1", "Control", 12.5, 2),
        ("C_CTRL_02", "DCS_RACK_L1", "COOLING_PUMP_P1", "Control", 14.0, 1),
        ("C_CTRL_03", "PLC_PANEL_L2", "MIXER_MOTOR_M2", "Control", 12.0, 1),
        ("C_CTRL_04", "PLC_PANEL_L2", "VALVE_CLUSTER_V1", "Control", 16.5, 2),
        ("C_SIG_01", "JUNC_L1_EAST", "TEMP_TRANSMITTER_TT101", "Signal", 8.8, 4),
        ("C_SIG_02", "JUNC_L2_WEST", "PRESSURE_SENSOR_PT202", "Signal", 9.2, 3),
        ("C_DATA_01", "DCS_RACK_L1", "PLC_PANEL_L2", "Data", 11.2, 2),
        ("C_DATA_02", "PLC_PANEL_L2", "REMOTE_IO_L3", "Data", 10.5, 1),
        ("C_BUS_01", "DCS_RACK_L1", "MCC_L1", "Bus", 9.5, 1),
        ("C_BUS_02", "PLC_PANEL_L2", "MCC_L2", "Bus", 9.5, 1),
    ]

    for r_idx, row in enumerate(sample_cables, 2):
        for c_idx, val in enumerate(row, 1):
            cell = ws_cables.cell(row=r_idx, column=c_idx, value=val)
            cell.font = regular_font
            cell.border = thin_border
            if c_idx in (1, 2, 3, 4):
                cell.alignment = left_align
            else:
                cell.alignment = right_align
        ws_cables.row_dimensions[r_idx].height = 20

    # Sheet 2: Branches
    ws_branches = wb.create_sheet(title="Branches")
    ws_branches.views.sheetView[0].showGridLines = True

    branch_headers = ["Tray ID", "From Node", "To Node", "Level", "Orientation", "Length (m)", "Tray Height (mm)"]
    for col_idx, h in enumerate(branch_headers, 1):
        c = ws_branches.cell(row=1, column=col_idx, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = center_align
        c.border = thin_border
    ws_branches.row_dimensions[1].height = 25

    sample_branches = [
        # Level 1 Horizontal
        ("BR_L1_01", "TRANSF_01", "MCC_L1", "Level 1", "horizontal", 12.0, 100.0),
        ("BR_L1_02", "MCC_L1", "DCS_RACK_L1", "Level 1", "horizontal", 8.5, 60.0),
        ("BR_L1_03", "MCC_L1", "JUNC_L1_EAST", "Level 1", "horizontal", 14.5, 60.0),
        ("BR_L1_04", "DCS_RACK_L1", "JUNC_L1_EAST", "Level 1", "horizontal", 10.0, 60.0),
        ("BR_L1_05", "JUNC_L1_EAST", "COOLING_PUMP_P1", "Level 1", "horizontal", 7.5, 60.0),
        ("BR_L1_06", "JUNC_L1_EAST", "TEMP_TRANSMITTER_TT101", "Level 1", "horizontal", 6.0, 60.0),
        ("BR_L1_07", "MCC_L1", "RISER_EAST_L1", "Level 1", "horizontal", 18.0, 100.0),
        ("BR_L1_08", "MCC_L1", "COMPRESSOR_CP1", "Level 1", "horizontal", 11.0, 60.0),
        # Riser 1 (Level 1 to Level 2)
        ("RISER_E_L1_L2", "RISER_EAST_L1", "RISER_EAST_L2", "Transition", "vertical", 5.0, 100.0),
        # Level 2 Horizontal
        ("BR_L2_01", "RISER_EAST_L2", "FEEDER_PANEL_L2", "Level 2", "horizontal", 15.0, 100.0),
        ("BR_L2_02", "FEEDER_PANEL_L2", "MCC_L2", "Level 2", "horizontal", 6.0, 100.0),
        ("BR_L2_03", "FEEDER_PANEL_L2", "PLC_PANEL_L2", "Level 2", "horizontal", 9.0, 60.0),
        ("BR_L2_04", "MCC_L2", "MIXER_MOTOR_M2", "Level 2", "horizontal", 12.5, 60.0),
        ("BR_L2_05", "PLC_PANEL_L2", "VALVE_CLUSTER_V1", "Level 2", "horizontal", 8.0, 60.0),
        ("BR_L2_06", "PLC_PANEL_L2", "JUNC_L2_WEST", "Level 2", "horizontal", 11.5, 60.0),
        ("BR_L2_07", "JUNC_L2_WEST", "PRESSURE_SENSOR_PT202", "Level 2", "horizontal", 5.0, 60.0),
        ("BR_L2_08", "RISER_EAST_L2", "RISER_EAST_L2_L3", "Level 2", "horizontal", 4.0, 100.0),
        # Riser 2 (Level 2 to Level 3)
        ("RISER_E_L2_L3", "RISER_EAST_L2_L3", "RISER_EAST_L3", "Transition", "vertical", 5.5, 100.0),
        # Level 3 Horizontal
        ("BR_L3_01", "RISER_EAST_L3", "REMOTE_IO_L3", "Level 3", "horizontal", 14.0, 60.0),
        ("BR_L3_02", "RISER_EAST_L3", "FOREHEARTH_FAN_M1", "Level 3", "horizontal", 22.0, 100.0),
        ("BR_L3_03", "RISER_EAST_L3", "EXHAUST_BLOWER_B1", "Level 3", "horizontal", 18.5, 100.0),
    ]

    for r_idx, row in enumerate(sample_branches, 2):
        for c_idx, val in enumerate(row, 1):
            cell = ws_branches.cell(row=r_idx, column=c_idx, value=val)
            cell.font = regular_font
            cell.border = thin_border
            if c_idx in (1, 2, 3, 4, 5):
                cell.alignment = left_align
            else:
                cell.alignment = right_align
        ws_branches.row_dimensions[r_idx].height = 20

    # Auto-adjust column widths
    for ws in [ws_cables, ws_branches]:
        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 5, 14)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
